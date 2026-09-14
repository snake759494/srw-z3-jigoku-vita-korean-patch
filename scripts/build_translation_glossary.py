#!/usr/bin/env python3
"""시옥편 번역 엑셀을 문맥형 시나리오 용어사전으로 통합한다.

원본 XLSX는 읽기만 하며 수정하지 않는다. 인물·로봇·공용 용어·rpw 문자열·
전투대사에서 일본어/한국어 쌍을 모아 하나의 JSON으로 만들고, 사람이 유지할
말투·관계·작품 분위기 규칙을 함께 기록한다. 이 사전은 자동 번역 결과를
만드는 용도가 아니라, 시나리오 문장을 검수할 때 표기와 화자 관계를 고정하는
공개 기준표다.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_EXCEL_ROOT = Path(r"D:\Z\psvita\시옥번역엑셀")
DEFAULT_OUTPUT = ROOT / "translations" / "retranslation" / "scenario_glossary_v2.json"
DEFAULT_SCHEMA = ROOT / "translations" / "retranslation" / "scenario_glossary_v2.schema.json"
SPACE_RE = re.compile(r"[ \t\u3000]+")
JAPANESE_RE = re.compile(r"[ぁ-ゖァ-ヺ一-龯々〆ヵヶ]")
PLACEHOLDER_RE = re.compile(r"^[\-ー―・･？?！!。．…]+$")


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def compact(value: Any) -> str:
    return SPACE_RE.sub(" ", text(value)).strip()


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_ref(path: Path, sheet: str, row_number: int) -> dict[str, Any]:
    return {"file": path.name, "sheet": sheet, "row": row_number}


def iter_table(path: Path, sheet_name: str, predicate) -> Iterable[tuple[int, list[str], list[Any]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    header: list[str] | None = None
    header_row = 0
    for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = [text(value) for value in row]
        if header is None:
            if predicate(values):
                header = values
                header_row = row_number
            continue
        if not any(values):
            continue
        yield row_number, header, list(row)
    wb.close()


def col_index(header: list[str], *names: str) -> int | None:
    for name in names:
        if name in header:
            return header.index(name)
    return None


def get(row: list[Any], header: list[str], *names: str) -> str:
    index = col_index(header, *names)
    if index is None or index >= len(row):
        return ""
    return text(row[index])


def is_real_term(ja: str, ko: str) -> bool:
    ja = compact(ja)
    ko = compact(ko)
    if not ja or not ko or PLACEHOLDER_RE.fullmatch(ja):
        return False
    if not JAPANESE_RE.search(ja) and not re.search(r"[A-Za-zＡ-Ｚａ-ｚ]", ja):
        return False
    if ko in {"-", "—", "?", "？", "없음"}:
        return False
    return True


def record_key(kind: str, ja: str) -> tuple[str, str]:
    return kind, compact(ja)


def add_record(records: dict[tuple[str, str], dict[str, Any]], *, kind: str, ja: str, ko: str, source: dict[str, Any], **extra: Any) -> None:
    ja_clean = compact(ja)
    ko_clean = compact(ko)
    if not is_real_term(ja_clean, ko_clean):
        return
    key = record_key(kind, ja_clean)
    item = records.setdefault(
        key,
        {
            "kind": kind,
            "ja": ja_clean,
            "preferredKo": ko_clean,
            "koCandidates": [],
            "aliasesKo": [],
            "series": [],
            "notes": [],
            "sources": [],
        },
    )
    if ko_clean not in item["koCandidates"]:
        item["koCandidates"].append(ko_clean)
    if ko_clean not in item["aliasesKo"] and ko_clean != item["preferredKo"]:
        item["aliasesKo"].append(ko_clean)
    if source not in item["sources"]:
        item["sources"].append(source)
    for key_name, value in extra.items():
        if value in (None, "", [], {}):
            continue
        if isinstance(value, list):
            target = item.setdefault(key_name, [])
            for member in value:
                if member not in target:
                    target.append(member)
        elif isinstance(value, dict):
            item.setdefault(key_name, {}).update(value)
        else:
            item[key_name] = value


def grouped_records(path: Path, sheet: str, predicate, group_key_names: tuple[str, ...]) -> dict[str, list[tuple[int, list[str], list[Any]]]]:
    groups: dict[str, list[tuple[int, list[str], list[Any]]]] = defaultdict(list)
    for row_number, header, row in iter_table(path, sheet, predicate):
        key = get(row, header, *group_key_names)
        if key:
            groups[key].append((row_number, header, row))
    return groups


def field_map(rows: list[tuple[int, list[str], list[Any]]]) -> dict[str, str]:
    result: dict[str, str] = {}
    for row_number, header, row in rows:
        tag = get(row, header, "필드태그")
        if tag:
            result[tag] = get(row, header, "일본어_원문")
            result[f"{tag}__ko"] = get(row, header, "한국어_수정")
            result[f"{tag}__meaning"] = get(row, header, "필드의미")
            result[f"{tag}__row"] = str(row_number)
    return result


def add_character_and_robot_data(root: Path, records: dict[tuple[str, str], dict[str, Any]], source_files: dict[str, Any]) -> dict[str, int]:
    counts = Counter()
    pt = root / "시옥편_MtZkn_Pt_최종파일_원문번역.xlsx"
    rt = root / "시옥편_MtZkn_Rt_최종파일_원문번역.xlsx"
    for path, kind, name_tags, group_key in [
        (pt, "character", ("CHFN", "CHNN"), ("Entry_ID", "내부파일명")),
        (rt, "robot", ("RBTN", "RBN2"), ("Entry_ID", "내부파일명")),
    ]:
        if not path.exists():
            continue
        source_files[path.name] = {"sha256": sha256(path), "role": kind}
        groups = grouped_records(path, "번역_데이터", lambda values: "필드태그" in values and "일본어_원문" in values, group_key)
        for group_id, rows in groups.items():
            fields = field_map(rows)
            series_ja, series_ko = fields.get("PRDC", ""), fields.get("PRDC__ko", "")
            voice_ja, voice_ko = fields.get("ACTR", ""), fields.get("ACTR__ko", "")
            description = compact(fields.get("DSCR__ko", ""))
            name_records = []
            for tag in name_tags:
                ja, ko = fields.get(tag, ""), fields.get(f"{tag}__ko", "")
                if not is_real_term(ja, ko):
                    continue
                extra = {
                    "series": [compact(series_ko)] if series_ko else [],
                    "seriesJa": [compact(series_ja)] if series_ja else [],
                    "profile": {"voiceActorJa": compact(voice_ja), "voiceActorKo": compact(voice_ko)} if voice_ja else {},
                    "descriptionHint": description[:600],
                    "entryId": group_id,
                }
                add_record(
                    records,
                    kind=kind,
                    ja=ja,
                    ko=ko,
                    source=source_ref(path, "번역_데이터", int(fields.get(f"{tag}__row", "0"))),
                    **extra,
                )
                name_records.append(ko)
                counts[kind] += 1
            # 파일럿 필드는 로봇과 별도로 캐릭터 표기 후보로 보관한다.
            pilot_ja, pilot_ko = fields.get("PLTN", ""), fields.get("PLTN__ko", "")
            if kind == "robot" and is_real_term(pilot_ja, pilot_ko):
                add_record(
                    records,
                    kind="pilot",
                    ja=pilot_ja,
                    ko=pilot_ko,
                    source=source_ref(path, "번역_데이터", int(fields.get("PLTN__row", "0"))),
                    series=[compact(series_ko)] if series_ko else [],
                    seriesJa=[compact(series_ja)] if series_ja else [],
                    associatedRobot=name_records,
                    entryId=group_id,
                )
                counts["pilot"] += 1
    return dict(counts)


def add_keyword_data(root: Path, records: dict[tuple[str, str], dict[str, Any]], source_files: dict[str, Any]) -> dict[str, int]:
    counts = Counter()
    # 두 파일은 같은 ``번역_데이터`` 시트 이름을 사용하지만 헤더 구조가
    # 다르다. 안내 행의 수에 의존하지 않고 헤더를 직접 찾는다.
    specs = [
        ("시옥편_MtZkn_KW_최종파일_원문번역.xlsx", "keyword_dictionary", "용어명_일본어 (WORD)"),
        ("시옥편_MtV_all_keyword_def_최종파일_원문번역.xlsx", "keyword_definition", "필드태그"),
    ]
    for filename, kind, header_marker in specs:
        path = root / filename
        if not path.exists():
            continue
        source_files[filename] = {"sha256": sha256(path), "role": kind}
        for row_number, header, row in iter_table(
            path,
            "번역_데이터",
            lambda values, marker=header_marker: marker in values,
        ):
            if header_marker == "용어명_일본어 (WORD)":
                ja = get(row, header, "용어명_일본어 (WORD)")
                ko = get(row, header, "용어명_한국어 (WORD)")
                description = get(row, header, "설명(현재)_한국어 (DSCR)")
                if is_real_term(ja, ko):
                    add_record(
                        records,
                        kind="term",
                        ja=ja,
                        ko=ko,
                        source=source_ref(path, "번역_데이터", row_number),
                        description=description[:1200],
                        entryId=get(row, header, "Entry_ID"),
                    )
                    counts["term"] += 1
            else:
                tag = get(row, header, "필드태그")
                if tag not in {"WORD", "SRCE"}:
                    continue
                ja = get(row, header, "일본어_원문")
                ko = get(row, header, "한국어_수정")
                if tag == "WORD" and is_real_term(ja, ko):
                    add_record(
                        records,
                        kind="term",
                        ja=ja,
                        ko=ko,
                        source=source_ref(path, "번역_데이터", row_number),
                        entryId=get(row, header, "External_ID"),
                    )
                    counts["term"] += 1
    return dict(counts)


def add_rpw_data(root: Path, records: dict[tuple[str, str], dict[str, Any]], source_files: dict[str, Any]) -> dict[str, int]:
    path = root / "시옥편_rpw_data_최종파일_한글데이터.xlsx"
    counts = Counter()
    if not path.exists():
        return {}
    source_files[path.name] = {"sha256": sha256(path), "role": "shared_string_pool"}
    for row_number, header, row in iter_table(path, "원문_번역_전체", lambda values: "일본어 원문" in values and "한국어 번역문" in values):
        ja = get(row, header, "일본어 원문")
        ko = get(row, header, "한국어 번역문")
        refs = get(row, header, "대응참조필드")
        if not is_real_term(ja, ko):
            continue
        categories = []
        if "pilot" in refs:
            categories.append("pilot_string")
        if "robot" in refs:
            categories.append("robot_string")
        if "weapon" in refs or "wpn" in refs:
            categories.append("weapon")
        if not categories:
            categories = ["shared_string"]
        for category in categories:
            add_record(records, kind=category, ja=ja, ko=ko, source=source_ref(path, "원문_번역_전체", row_number), references=refs[:1200])
            counts[category] += 1
    return dict(counts)


def add_scenario_and_battle_context(root: Path, source_files: dict[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {"scenarioSummaries": [], "battleStyle": {"masterLines": 0, "variantRows": 0, "samples": []}}
    kdata = root / "시옥편_KDataVITA_최종파일_원문번역.xlsx"
    if kdata.exists():
        source_files[kdata.name] = {"sha256": sha256(kdata), "role": "scenario_summary"}
        for sheet, title_cols, body_cols in [
            ("본편_시나리오", ("일본어 원문", "한국어 번역"), ("일본어 원문", "한국어 번역")),
            ("DLC_시나리오", ("일본어 제목", "한국어 제목"), ("일본어 본문", "한국어 본문")),
        ]:
            for row_number, header, row in iter_table(kdata, sheet, lambda values, cols=title_cols: all(col in values for col in cols)):
                if sheet == "본편_시나리오":
                    ja, ko = get(row, header, "일본어 원문"), get(row, header, "한국어 번역")
                    record_id = get(row, header, "레코드 ID")
                else:
                    ja, ko = get(row, header, "일본어 제목") + "\n" + get(row, header, "일본어 본문"), get(row, header, "한국어 제목") + "\n" + get(row, header, "한국어 본문")
                    record_id = get(row, header, "레코드 ID")
                if ja and ko:
                    result["scenarioSummaries"].append({"asset": record_id, "ja": ja, "koReference": ko, "source": source_ref(kdata, sheet, row_number)})
    srvc = root / "시옥편_SRVC_최종파일_원문번역.xlsx"
    if srvc.exists():
        source_files[srvc.name] = {"sha256": sha256(srvc), "role": "battle_dialogue"}
        for row_number, header, row in iter_table(srvc, "전체 전투대사", lambda values: "일본어 원문" in values and "최종 한국어" in values):
            ja, ko = get(row, header, "일본어 원문"), get(row, header, "최종 한국어")
            if not ja or not ko:
                continue
            result["battleStyle"]["variantRows"] += 1
            if get(row, header, "Variant Key").endswith("-1"):
                result["battleStyle"]["masterLines"] += 1
            if len(result["battleStyle"]["samples"]) < 80 and ("！" in ja or "!" in ja or "ぞ" in ja or "ろ" in ja):
                result["battleStyle"]["samples"].append({"ja": ja, "ko": ko, "source": source_ref(srvc, "전체 전투대사", row_number)})
    return result


SPEECH_PROFILES = [
    {"id": "hibiki_kujo", "namesJa": ["ヒビキ", "ヒビキ・カミシロ"], "namesKo": ["히비키", "히비키 카미시로"], "series": ["오리지널"], "style": "주인공. 기본은 동료에게 반말과 단정한 말투를 섞고, 위기에서는 짧고 빠르게 끊는다.", "honorifics": "스즈네에게는 선생님이라는 관계를 의식해 존댓말과 반말이 흔들릴 수 있으나, 감정이 드러날수록 반말이 강해진다.", "relations": ["suzune_anemone", "z_blue"], "evidence": "시옥편 오리지널 주인공 설정과 시나리오 화자 흐름"},
    {"id": "suzune_anemone", "namesJa": ["スズネ", "スズネ・アネモネ"], "namesKo": ["스즈네", "스즈네 아네모네"], "series": ["오리지널"], "style": "교사·항법사. 평상시에는 학생을 다독이는 존댓말, 당황하거나 본성이 드러날 때는 거친 반말과 감정적 표현.", "honorifics": "히비키를 학생으로 부르며 공적인 장면에서는 존댓말을 유지한다.", "relations": ["hibiki_kujo"], "evidence": "KDataVITA 시나리오 요약·Pt 인물 설명"},
    {"id": "shin_asuka", "namesJa": ["シン", "シン・アスカ"], "namesKo": ["신", "신 아스카"], "series": ["기동전사 건담 SEED DESTINY"], "style": "직선적이고 감정적인 반말. 분노·상실·정의감이 앞설 때 문장을 짧게 끊고 목소리가 높아진다.", "honorifics": "키라·아스란에게는 관계와 장면에 따라 반말을 쓰되, 공식 지휘 상황에서는 군대식 단정체로 낮춘다.", "relations": ["kira_yamato", "aslan_zala"], "evidence": "Pt 인물명·작품 표기, 전투대사 말투 표본"},
    {"id": "kira_yamato", "namesJa": ["キラ", "キラ・ヤマト"], "namesKo": ["키라", "키라 야마토"], "series": ["기동전사 건담 SEED DESTINY"], "style": "온화하고 부드러운 반말. 상대를 몰아붙이기보다 설득하며, 결심한 순간에는 짧고 단정한 문장으로 바뀐다.", "honorifics": "아스란·신에게는 친밀한 반말, 상관·공식 보고에는 차분한 존댓말 또는 단정체.", "relations": ["shin_asuka", "aslan_zala"], "evidence": "Pt 인물명·작품 표기, 전투대사 말투 표본"},
    {"id": "aslan_zala", "namesJa": ["アスラン", "アスラン・ザラ"], "namesKo": ["아스란", "아스란 자라"], "series": ["기동전사 건담 SEED DESTINY"], "style": "책임감 있는 군인 말투. 감정이 흔들려도 명령·판단은 단정하게 말한다.", "honorifics": "키라와 신에게는 친우로서 반말, 상관과 공식 회의에서는 존댓말.", "relations": ["kira_yamato", "shin_asuka"], "evidence": "Pt 인물명·작품 표기"},
    {"id": "kamille_bidan", "namesJa": ["カミーユ", "カミーユ・ビダン"], "namesKo": ["카미유", "카미유 비단"], "series": ["기동전사 Z건담"], "style": "예민하고 지적인 청년. 동료에게는 반말, 윗사람에게는 존댓말을 쓰되 분노하면 직설적으로 변한다.", "honorifics": "아무로·크와트로 등 선배에게는 기본적으로 존중을 유지한다.", "relations": ["amuro_ray", "quattro_bajeena"], "evidence": "Pt 인물명·작품 표기"},
    {"id": "amuro_ray", "namesJa": ["アムロ", "アムロ・レイ"], "namesKo": ["아무로", "아무로 레이"], "series": ["기동전사 건담 역습의 샤아"], "style": "침착한 베테랑. 감정을 과장하지 않고 군인다운 단정체로 핵심만 말한다.", "honorifics": "카미유 등 후배에게는 부드러운 반말, 공식 보고나 상관에게는 존댓말.", "relations": ["kamille_bidan", "char_aznable"], "evidence": "Pt 인물명·작품 표기"},
    {"id": "char_aznable", "namesJa": ["シャア", "シャア・アズナブル", "クワトロ", "クワトロ・バジーナ"], "namesKo": ["샤아", "샤아 아즈나블", "크와트로", "크와트로 바지나"], "series": ["기동전사 건담 역습의 샤아", "기동전사 Z건담"], "style": "카리스마와 냉소가 섞인 지휘관 말투. 짧은 명령형과 비유를 사용하며 정체·감정을 한 번에 드러내지 않는다.", "honorifics": "공식 지휘에서는 존칭보다 직위와 명령을 우선한다.", "relations": ["amuro_ray", "kamille_bidan"], "evidence": "Pt 인물명·작품 표기"},
    {"id": "setsuna_f_seiei", "namesJa": ["刹那", "刹那・F・セイエイ"], "namesKo": ["세츠나", "세츠나 F 세이에이"], "series": ["기동전사 건담 00"], "style": "말수가 적고 목적 지향적인 반말. 군더더기를 줄이고 선언형 문장을 쓴다.", "honorifics": "동료에게는 반말, 작전 보고에서는 짧은 단정체.", "relations": ["heero_yuy", "lockon_stratos"], "evidence": "Pt 인물명·작품 표기"},
    {"id": "heero_yuy", "namesJa": ["ヒイロ", "ヒイロ・ユイ"], "namesKo": ["히이로", "히이로 유이"], "series": ["신기동전기 건담 W"], "style": "감정을 억제한 짧은 문장. 불필요한 주어·수식을 생략하고 임무를 우선한다.", "honorifics": "대체로 반말 또는 명령형이며, 관계보다 임무가 앞선다.", "relations": ["setsuna_f_seiei"], "evidence": "Pt 인물명·작품 표기"},
    {"id": "roger_smith", "namesJa": ["ロジャー", "ロジャー・スミス"], "namesKo": ["로저", "로저 스미스"], "series": ["THE 빅오"], "style": "지적이고 무게감 있는 반말. 협상가답게 완곡하지만 결론은 분명하다.", "honorifics": "도로시에게는 건조한 반말, 낯선 상대에게는 예의를 갖춘 존댓말.", "relations": ["r_dorothy_wayneright"], "evidence": "Pt 인물명·작품 표기"},
    {"id": "r_dorothy_wayneright", "namesJa": ["ドロシー", "Ｒ・ドロシー・ウェインライト"], "namesKo": ["도로시", "R. 도로시 웨인라이트"], "series": ["THE 빅오"], "style": "짧고 담담한 반말. 감정 표현을 과장하지 않고 관찰한 사실을 건조하게 말한다.", "honorifics": "로저에게도 존댓말보다 건조한 반말을 유지한다.", "relations": ["roger_smith"], "evidence": "Pt 인물명·작품 표기"},
    {"id": "sousuke_sagara", "namesJa": ["宗介", "相良宗介"], "namesKo": ["소스케", "사가라 소스케"], "series": ["풀 메탈 패닉!"], "style": "군인식 단정체와 과도하게 진지한 직역투가 코미디를 만든다. 전투에서는 명령형, 학교에서는 상황을 오해한 딱딱한 말투.", "honorifics": "카나메에게는 이름을 부르며 반말, 상관에게는 군대식 존댓말.", "relations": ["kaname_chidori", "mao", "kurz_weber"], "evidence": "KDataVITA 시나리오 요약·전투대사"},
    {"id": "kaname_chidori", "namesJa": ["かなめ", "千鳥かなめ"], "namesKo": ["카나메", "치도리 카나메"], "series": ["풀 메탈 패닉!"], "style": "현실적인 고등학생 말투. 빠른 반응, 핀잔, 츳코미를 살리고 감정이 격해지면 반말을 강하게 한다.", "honorifics": "소스케에게는 반말, 교사·상관에게는 상황에 맞춰 예의를 갖춘다.", "relations": ["sousuke_sagara"], "evidence": "KDataVITA 시나리오 요약"},
    {"id": "alto_saotome", "namesJa": ["アルト", "早乙女アルト"], "namesKo": ["알토", "사오토메 알토"], "series": ["마크로스 프론티어"], "style": "자존심이 강하고 직설적인 반말. 무대·전투에서 모두 과장된 선언을 쓰지만 감정의 핵심은 솔직하다.", "honorifics": "셰릴·란카에게는 관계에 따른 반말, 어른에게는 불만이 있어도 기본 예의를 남긴다.", "relations": ["sheryl_nome", "ranka_lee"], "evidence": "KDataVITA 시나리오 요약·전투대사"},
    {"id": "sheryl_nome", "namesJa": ["シェリル", "シェリル・ノーム"], "namesKo": ["셰릴", "셰릴 놈"], "series": ["마크로스 프론티어"], "style": "자신감 있고 화려한 말투. 무대에서는 당당하고, 사적인 장면에서는 외로움과 약함이 드러난다.", "honorifics": "알토에게는 장난기 섞인 반말, 공식 자리에서는 스타다운 품위를 유지한다.", "relations": ["alto_saotome", "ranka_lee"], "evidence": "전투대사·시나리오 요약"},
    {"id": "ranka_lee", "namesJa": ["ランカ", "ランカ・リー"], "namesKo": ["란카", "란카 리"], "series": ["마크로스 프론티어"], "style": "밝고 순수한 존댓말·반말 혼용. 감정이 앞설 때 말이 빠르고 솔직해진다.", "honorifics": "알토·셰릴에게는 친밀도에 따른 말투를 유지한다.", "relations": ["alto_saotome", "sheryl_nome"], "evidence": "전투대사·시나리오 요약"},
]


RELATIONSHIPS = [
    {"from": "hibiki_kujo", "to": "suzune_anemone", "relation": "학생과 교사·항법사", "speech": "히비키는 스즈네를 기본적으로 선생님으로 존중하지만 감정이 격해질수록 반말이 섞인다.", "priority": "scene"},
    {"from": "sousuke_sagara", "to": "kaname_chidori", "relation": "경호 대상과 경호원·동급생", "speech": "카나메는 소스케를 거칠게 타박하고, 소스케는 군인식으로 진지하게 응답한다.", "priority": "scene"},
    {"from": "kira_yamato", "to": "shin_asuka", "relation": "전쟁을 겪은 선후배이자 동료", "speech": "키라는 부드럽게 설득하고 신은 먼저 감정적으로 반발하지만, 전투에서는 서로를 동료로 인정한다.", "priority": "scene"},
    {"from": "kira_yamato", "to": "aslan_zala", "relation": "오랜 친구이자 전우", "speech": "친밀한 반말을 유지하되 공식 지휘 장면에서는 군인다운 단정체로 전환한다.", "priority": "scene"},
    {"from": "amuro_ray", "to": "kamille_bidan", "relation": "선배 뉴타입과 후배 파일럿", "speech": "아무로는 카미유를 다그치기보다 안정시키며, 카미유는 기본적으로 예의를 남긴다.", "priority": "scene"},
    {"from": "roger_smith", "to": "r_dorothy_wayneright", "relation": "협상가와 안드로이드 동료", "speech": "로저는 정중한 어휘를 쓰지만 말투는 무겁고, 도로시는 짧고 건조하게 받아친다.", "priority": "scene"},
    {"from": "alto_saotome", "to": "sheryl_nome", "relation": "파일럿과 가수·라이벌", "speech": "셰릴은 도발적이고 당당하며 알토는 자존심 때문에 반발하지만 서로를 의식한다.", "priority": "scene"},
    {"from": "alto_saotome", "to": "ranka_lee", "relation": "파일럿과 후배 가수·동료", "speech": "란카의 순수한 존댓말과 알토의 투박한 반말 대비를 살린다.", "priority": "scene"},
]


SERIES_STYLES = [
    {"ja": "機動戦士ガンダムSEED DESTINY", "ko": "기동전사 건담 SEED DESTINY", "tone": "전쟁 후유증·상실·책임의 대화. 군사 지휘와 개인 감정의 간극을 분명히 한다."},
    {"ja": "機動戦士ガンダム００", "ko": "기동전사 건담 00", "tone": "짧은 선언형 문장과 조직·임무의 어휘. 세츠나는 감정을 설명하지 않고 결론을 말한다."},
    {"ja": "機動戦士Ｚガンダム", "ko": "기동전사 Z건담", "tone": "군·조직의 위계와 개인의 예민함을 함께 살린다. 선후배 관계의 존칭을 고정한다."},
    {"ja": "フルメタル・パニック！", "ko": "풀 메탈 패닉!", "tone": "군사 전문용어를 정확히 쓰되 학교 코미디 장면의 오해와 츳코미를 죽이지 않는다."},
    {"ja": "マクロスＦ", "ko": "마크로스 프론티어", "tone": "무대의 화려함, 삼각관계의 미묘함, 조종사의 자존심을 대사 리듬으로 구분한다."},
    {"ja": "創聖のアクエリオン／アクエリオンＥＶＯＬ", "ko": "창성의 아쿠에리온／아쿠에리온 EVOL", "tone": "신화적 표현과 청춘의 감정을 함께 살린다. 합체·운명 어휘는 장엄하게 처리한다."},
    {"ja": "ＴＨＥビッグオー", "ko": "THE 빅오", "tone": "느리고 무게감 있는 협상가의 문장과 도로시의 건조한 응답을 대비한다."},
    {"ja": "天元突破グレンラガン", "ko": "천원돌파 그렌라간", "tone": "과장된 선언과 동료를 끌어올리는 열기를 살리되 반복되는 구호는 과도하게 번역하지 않는다."},
    {"ja": "鉄人２８号", "ko": "철인 28호", "tone": "소년 탐정극의 명료한 문장과 어른의 책임감을 구분한다."},
    {"ja": "無敵ロボ　トライダーＧ７", "ko": "무적로보 트라이더 G7", "tone": "초등학생 사장 왓타의 장난기와 회사원들의 생활 코미디를 함께 살린다."},
]


def build(args: argparse.Namespace) -> dict[str, Any]:
    root = args.excel_root.expanduser().resolve()
    records: dict[tuple[str, str], dict[str, Any]] = {}
    source_files: dict[str, Any] = {}
    extraction_counts: dict[str, Any] = {}
    extraction_counts["names"] = add_character_and_robot_data(root, records, source_files)
    extraction_counts["keywords"] = add_keyword_data(root, records, source_files)
    extraction_counts["rpw"] = add_rpw_data(root, records, source_files)
    context = add_scenario_and_battle_context(root, source_files)

    # 용어·인명 표기는 기존 프로젝트 기준을 우선하되, 엑셀의 모든 후보를
    # aliasesKo에 남겨 사람이 차이를 검토할 수 있게 한다.
    manual_canonical = {
        "時獄戦役": "시옥전쟁",
        "時獄篇": "시옥편",
        "破界事変": "파계사변",
        "再世戦争": "재세전쟁",
        "大時空震動": "대시공진동",
        "大時空震": "대시공진동",
        "時空震動": "시공진동",
        "時空震": "시공진동",
        "次元震": "차원진동",
        "次元の穴": "차원의 구멍 《어비스》",
        "アビス": "어비스",
        "烙印": "낙인(스티그마)",
        "スフィア": "스피어",
        "インサラウム": "인사라움",
        "ＺＥＵＴＨ": "ZEUTH",
        "ＺＥＸＩＳ": "ZEXIS",
        "ＵＣＷ": "UCW",
        "ＡＤＷ": "ADW",
    }
    for key, preferred in manual_canonical.items():
        for kind in ("term", "keyword_dictionary", "keyword_definition", "shared_string"):
            item = records.get((kind, key))
            if item:
                if item["preferredKo"] != preferred and item["preferredKo"] not in item["aliasesKo"]:
                    item["aliasesKo"].append(item["preferredKo"])
                item["preferredKo"] = preferred

    for item in records.values():
        item["koCandidates"] = list(dict.fromkeys(item["koCandidates"]))
        item["aliasesKo"] = [value for value in item["aliasesKo"] if value != item["preferredKo"]]
        item["sources"] = sorted(item["sources"], key=lambda ref: (ref["file"], ref["sheet"], ref["row"]))
        if item["kind"] in {"character", "pilot"}:
            item.setdefault("usage", "인물명은 CHFN/CHNN·PLTN 표기를 우선하며, 대사에서는 화자 관계에 따라 호칭을 조절한다.")
        elif item["kind"] in {"robot", "robot_string", "weapon"}:
            item.setdefault("usage", "기체·무기명은 전투·메뉴·사전에서 동일 표기를 사용하고, 설명문에서는 자연스러운 조사를 붙인다.")
        elif item["kind"] == "term":
            item.setdefault("usage", "시나리오·사전·전투대사에서 같은 개념은 같은 한국어 표기를 유지한다.")

    # 기본 프로젝트 용어집의 표기와 인물 프로필을 합쳐 누락된 고정명을 보충한다.
    existing_glossary = ROOT / "translations" / "retranslation" / "glossary.json"
    if existing_glossary.exists():
        existing = json.loads(existing_glossary.read_text(encoding="utf-8"))
        for key, value in existing.get("terms", {}).items():
            if isinstance(value, str):
                add_record(records, kind="project_term", ja=key, ko=value, source={"file": existing_glossary.name, "sheet": "terms", "row": key})

    return {
        "$schema": "./scenario_glossary_v2.schema.json",
        "format": "siok.scenario-glossary",
        "formatVersion": 2,
        "generatedAt": "2026-08-12",
        "sourcePolicy": {
            "translationInput": "시나리오 일본어 원문",
            "referenceUse": "엑셀의 한국어는 고유명사·관계·표기·품질 비교를 위한 참고값이며 문장 자동 복사의 근거가 아니다.",
            "originalGame": "본인 소유 게임에서 추출한 자료만 사용하며 게임 원본·CPK·EBOOT는 저장소에 포함하지 않는다.",
            "priority": ["수동 canonical 표기", "엑셀 원문-수정 쌍", "프로젝트 기존 용어집", "장면별 관계"],
        },
        "sourceFolder": root.name,
        "sourceFiles": source_files,
        "extractionCounts": extraction_counts,
        "records": sorted(records.values(), key=lambda item: (item["kind"], item["ja"])),
        "speechProfiles": SPEECH_PROFILES,
        "relationships": RELATIONSHIPS,
        "seriesStyles": SERIES_STYLES,
        "context": {
            "worlds": [
                {"ja": "ＵＣＷ", "ko": "UCW", "note": "ZEUTH가 원래 있던 세계권."},
                {"ja": "ＡＤＷ", "ko": "ADW", "note": "파계사변·재세전쟁·시옥전쟁이 이어지는 세계권."},
            ],
            "wars": [
                {"ja": "破界事変", "ko": "파계사변", "note": "제2차 Z 파계편의 차원 간 전쟁."},
                {"ja": "再世戦争", "ko": "재세전쟁", "note": "파계사변 약 1년 후의 전쟁."},
                {"ja": "時獄戦役", "ko": "시옥전쟁", "note": "제3차 Z 시옥편 본편의 전쟁."},
            ],
            "factions": [
                {"ja": "ＺＥＵＴＨ", "ko": "ZEUTH", "role": "UCW 출신의 전투 집단."},
                {"ja": "ＺＥＸＩＳ", "ko": "ZEXIS", "role": "파계사변·재세전쟁을 거친 연합 전력."},
                {"ja": "Ｚ－ＢＬＵＥ", "ko": "Z-BLUE", "role": "시옥전쟁에서 결성되는 주인공 측 연합."},
                {"ja": "ミスリル", "ko": "미스릴", "role": "풀 메탈 패닉의 특수부대. 군사 전문용어와 계급을 유지한다."},
                {"ja": "Ｓ．Ｍ．Ｓ", "ko": "S.M.S", "role": "마크로스 프론티어의 민간 군사조직."},
                {"ja": "インサラウム", "ko": "인사라움", "role": "재세전쟁과 시옥편을 잇는 다원세계 왕국."},
            ],
        },
        "scenarioSummaries": context["scenarioSummaries"],
        "battleStyle": context["battleStyle"],
        "referenceRulesFile": "reference_rules.json",
        "seriesReferenceFile": "series_reference.json",
        "referenceReview": {
            "checkedAt": "2026-08-14",
            "sourceFolder": "D:\\Z\\psvita\\시옥번역엑셀",
            "workCount": 32,
            "note": "엑셀에서 재생성한 고유명사·용어·공용 문자열은 원문 표기의 기준으로 사용하고, 작품별 관계·말투는 series_reference.json을 보조 기준으로 사용한다.",
        },
        "qualityRules": [
            "이름 행은 사전의 preferredKo를 사용하고, 문장 속 호칭은 관계 규칙을 우선한다.",
            "신·키라·카미유·아무로·세츠나·히이로·로저·도로시의 말투를 한 문장씩 번갈아 섞지 않는다.",
            "선생님·상관·선배·동료의 존칭은 장면의 화자 관계를 기준으로 일관되게 유지한다.",
            "원문 행 경계와 실제 줄바꿈은 유지하고, 대사창 byteLimit/encodedLength를 넘기면 문장을 압축한다.",
            "⑲·⑳·㊥·㊦·㊧·㊨·$n·$l·$c 같은 제어문자의 종류·개수·순서를 바꾸지 않는다.",
            "번역 완료 전 상태는 draft이며, 사람이 문맥·원작 느낌·고유명사·길이를 확인한 뒤 reviewed로 올린다.",
        ],
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel-root", type=Path, default=DEFAULT_EXCEL_ROOT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--schema-output", type=Path, default=DEFAULT_SCHEMA)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    result = build(args)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    schema = {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "$id": "urn:siok:scenario-glossary:v2",
        "type": "object",
        "required": ["format", "formatVersion", "records", "speechProfiles", "relationships", "seriesStyles", "qualityRules"],
        "properties": {
            "format": {"const": "siok.scenario-glossary"},
            "formatVersion": {"const": 2},
            "records": {"type": "array", "items": {"type": "object", "required": ["kind", "ja", "preferredKo", "sources"]}},
            "speechProfiles": {"type": "array"},
            "relationships": {"type": "array"},
            "seriesStyles": {"type": "array"},
            "qualityRules": {"type": "array", "minItems": 1},
        },
        "additionalProperties": True,
    }
    args.schema_output.parent.mkdir(parents=True, exist_ok=True)
    args.schema_output.write_text(json.dumps(schema, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), "records": len(result["records"]), "sources": len(result["sourceFiles"]), "counts": result["extractionCounts"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
